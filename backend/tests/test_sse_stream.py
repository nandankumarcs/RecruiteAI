"""Tests for SSE streaming endpoint."""

from __future__ import annotations

import asyncio
import json
from io import BytesIO

import pytest
import pytest_asyncio
from docx import Document
from httpx import AsyncClient

from app.agents.resume_parser_agent import get_resume_parser_agent
from app.main import app
from app.models.job import Job
from app.services.resume_session_manager import get_session_manager


class FakeResumeParserAgent:
    """Deterministic parser for testing."""

    async def parse_resume(self, file_path: str, file_type: str) -> dict:
        return {
            "candidate_name": "Test Candidate",
            "phone_number": "+1 111 222 3333",
            "email": "test@example.com",
            "raw_text": "Test Candidate\ntest@example.com\n+1 111 222 3333",
            "parsed_data": {"schema_version": "resume.v2"},
        }

    async def evaluate_candidate_against_jd(self, raw_text: str, job_description: str):
        from dataclasses import dataclass
        
        @dataclass
        class Evaluation:
            matching_score: float
            explanation: str
        
        return Evaluation(matching_score=85.0, explanation="Good match")


@pytest_asyncio.fixture
async def job_for_sse(db_session, test_user) -> Job:
    """Create a job for SSE testing."""
    job = Job(
        user_id=test_user.id,
        title="SSE Test Job",
        description="Testing SSE streaming",
        requirements="Python",
        status="active",
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)
    return job


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Generate an in-memory DOCX resume."""
    document = Document()
    document.add_paragraph("Test Candidate")
    document.add_paragraph("test@example.com")
    document.add_paragraph("+1 111 222 3333")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest_asyncio.fixture(autouse=True)
async def override_parser():
    """Use fake parser for tests."""
    app.dependency_overrides[get_resume_parser_agent] = lambda: FakeResumeParserAgent()
    yield
    app.dependency_overrides.pop(get_resume_parser_agent, None)


@pytest.mark.asyncio
async def test_sse_stream_endpoint_exists(
    authenticated_client: AsyncClient, job_for_sse: Job
):
    """Test that SSE endpoint exists and requires valid session."""
    # Try to connect without valid session
    response = await authenticated_client.get(
        f"/api/jobs/{job_for_sse.id}/resumes/stream?session_id=invalid-session",
        timeout=5.0
    )
    
    # Should return 404 for invalid session
    assert response.status_code == 404


@pytest.mark.asyncio
async def test_sse_stream_validates_job_ownership(
    authenticated_client: AsyncClient, db_session, test_user
):
    """Test that SSE endpoint validates job ownership."""
    # Create a job for another user
    from app.models.user import User
    from app.core.security import hash_password
    import uuid
    
    other_user = User(
        id=uuid.uuid4(),
        email="other@example.com",
        hashed_password=hash_password("Password123"),
        full_name="Other User",
        is_active=True,
    )
    db_session.add(other_user)
    await db_session.commit()
    
    other_job = Job(
        user_id=other_user.id,
        title="Other Job",
        description="Not mine",
        requirements="None",
        status="active",
    )
    db_session.add(other_job)
    await db_session.commit()
    await db_session.refresh(other_job)
    
    # Create a session for the other job
    session_manager = get_session_manager()
    session_id = session_manager.create_session(
        job_id=str(other_job.id),
        user_id=str(other_user.id),
        total_resumes=1
    )
    
    # Try to access with authenticated_client (test_user)
    response = await authenticated_client.get(
        f"/api/jobs/{other_job.id}/resumes/stream?session_id={session_id}",
        timeout=5.0
    )
    
    # Should return 404 because job doesn't belong to test_user
    assert response.status_code == 404


@pytest.mark.asyncio
async def test_upload_returns_session_id(
    authenticated_client: AsyncClient, job_for_sse: Job, sample_docx_bytes: bytes
):
    """Test that upload endpoint returns session_id."""
    response = await authenticated_client.post(
        f"/api/jobs/{job_for_sse.id}/resumes",
        files=[("files", ("test.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    
    assert response.status_code == 202
    data = response.json()
    assert "session_id" in data
    assert "total_resumes" in data
    assert data["total_resumes"] == 1


@pytest.mark.asyncio
async def test_sse_stream_basic_flow(
    authenticated_client: AsyncClient, job_for_sse: Job, sample_docx_bytes: bytes
):
    """Test basic SSE streaming flow."""
    # Upload resume to get session_id
    upload_response = await authenticated_client.post(
        f"/api/jobs/{job_for_sse.id}/resumes",
        files=[("files", ("test.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    
    assert upload_response.status_code == 202
    session_id = upload_response.json()["session_id"]
    
    # Connect to SSE stream
    events_received = []
    
    async with authenticated_client.stream(
        "GET",
        f"/api/jobs/{job_for_sse.id}/resumes/stream?session_id={session_id}",
        timeout=10.0
    ) as response:
        assert response.status_code == 200
        assert response.headers["content-type"] == "text/event-stream; charset=utf-8"
        
        # Read events until all_completed or timeout
        async for line in response.aiter_lines():
            if line.startswith("event:"):
                event_type = line.split(":", 1)[1].strip()
            elif line.startswith("data:"):
                data = json.loads(line.split(":", 1)[1].strip())
                events_received.append({"type": event_type, "data": data})
                
                # Stop after all_completed
                if event_type == "all_completed":
                    break
    
    # Verify we received expected events
    event_types = [e["type"] for e in events_received]
    
    # Should have: starting, uploading, parsing, evaluating, completed, all_completed
    assert "starting" in event_types
    assert "uploading" in event_types
    assert "parsing" in event_types
    assert "evaluating" in event_types
    assert "completed" in event_types
    assert "all_completed" in event_types
    
    # Verify event data structure
    for event in events_received:
        if event["type"] != "all_completed":
            assert "session_id" in event["data"]
            assert "resume_index" in event["data"]
            assert "total_resumes" in event["data"]
            assert "filename" in event["data"]
            assert "timestamp" in event["data"]
    
    # Verify completed event has candidate details
    completed_events = [e for e in events_received if e["type"] == "completed"]
    assert len(completed_events) == 1
    assert completed_events[0]["data"]["candidate_name"] == "Test Candidate"
    assert completed_events[0]["data"]["email"] == "test@example.com"
    assert completed_events[0]["data"]["matching_score"] == 85.0
