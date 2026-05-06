"""Tests for the ResumeParserAgent text extraction and structured parsing."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.agents.resume_parser_agent import (
    CertificationEntry,
    EducationEntry,
    ExperienceEntry,
    ProjectEntry,
    ResumeParserAgent,
    SkillCategory,
    StructuredResumeData,
)


def create_sample_docx(file_path: Path):
    """Create a realistic DOCX resume fixture on disk."""
    document = Document()
    document.add_paragraph("John Doe")
    document.add_paragraph("john.doe@example.com")
    document.add_paragraph("+1 (555) 123-4567")
    document.add_paragraph("Skills")
    document.add_paragraph("Python")
    document.add_paragraph("FastAPI")
    document.add_paragraph("Experience")
    document.add_paragraph("Senior Engineer Aug 2022 – Jan 2024")
    document.add_paragraph("Acme Corp New Delhi, India")
    document.add_paragraph("• Built backend APIs")
    document.add_paragraph("Projects")
    document.add_paragraph("Candidate Platform")
    document.add_paragraph("Python • FastAPI • PostgreSQL")
    document.add_paragraph("• Built resume parsing workflows")
    document.add_paragraph("Education")
    document.add_paragraph("University of Testing Jul 2018 – May 2022")
    document.add_paragraph("B.Tech Computer Science — CGPA: 8.8/10 Noida, India")
    document.add_paragraph("Technical Skills")
    document.add_paragraph("Languages: Python, SQL")
    document.add_paragraph("Backend: FastAPI, PostgreSQL")
    document.add_paragraph("Certifications")
    document.add_paragraph("Cloud Basics – Example Academy (2024)")
    document.save(file_path)


def test_extract_text_from_docx(tmp_path: Path):
    """DOCX extraction should preserve key lines from the document."""
    file_path = tmp_path / "resume.docx"
    create_sample_docx(file_path)

    agent = ResumeParserAgent(enable_llm=False)
    text = agent.extract_text_from_docx(str(file_path))

    assert "John Doe" in text
    assert "john.doe@example.com" in text
    assert "Senior Engineer Aug 2022 – Jan 2024" in text


def test_extract_phone_and_email():
    """Regex extraction should capture first contact details."""
    raw_text = """
    John Doe
    john.doe@example.com
    +1 (555) 123-4567
    """
    agent = ResumeParserAgent(enable_llm=False)

    assert agent.extract_email(raw_text) == "john.doe@example.com"
    assert agent.extract_phone_number(raw_text) == "+1 (555) 123-4567"


@pytest.mark.asyncio
async def test_parse_resume_uses_structured_data(monkeypatch, tmp_path: Path):
    """Parsed resume output should combine text extraction and structured fields."""
    file_path = tmp_path / "resume.docx"
    create_sample_docx(file_path)
    agent = ResumeParserAgent(enable_llm=False)

    async def fake_extract_structured_data(_: str):
        return StructuredResumeData(
            candidate_name="John Doe",
            email="john.doe@example.com",
            skills=["Python", "FastAPI"],
            skill_categories=[SkillCategory(category="Backend", items=["Python", "FastAPI"])],
            experience=[
                ExperienceEntry(
                    role_title="Senior Engineer",
                    company_name="Acme Corp",
                    from_date="Aug 2022",
                    to_date="Jan 2024",
                    tasks_performed=["Built backend APIs"],
                )
            ],
            projects=[
                ProjectEntry(
                    name="Candidate Platform",
                    technologies=["Python", "FastAPI"],
                    bullets=["Built resume parsing workflows"],
                )
            ],
            education=[
                EducationEntry(
                    institution="University of Testing",
                    degree="B.Tech Computer Science",
                    score="CGPA: 8.8/10",
                )
            ],
            certifications=[CertificationEntry(name="Cloud Basics", issuer="Example Academy", year="2024")],
            summary="Experienced backend engineer.",
        )

    monkeypatch.setattr(agent, "extract_structured_data", fake_extract_structured_data)

    parsed = await agent.parse_resume(str(file_path), "docx")

    assert parsed["candidate_name"] == "John Doe"
    assert parsed["email"] == "john.doe@example.com"
    assert parsed["phone_number"] == "+1 (555) 123-4567"
    assert parsed["parsed_data"]["skills"] == ["Python", "FastAPI"]
    assert parsed["parsed_data"]["schema_version"] == "resume.v2"
    assert parsed["parsed_data"]["summary"] == "Experienced backend engineer."
    assert parsed["parsed_data"]["experience"][0]["role_title"] == "Senior Engineer"
    assert parsed["parsed_data"]["experience"][0]["company_name"] == "Acme Corp"
    assert parsed["parsed_data"]["experience"][0]["from_date"] == "Aug 2022"
    assert parsed["parsed_data"]["experience"][0]["tasks_performed"] == ["Built backend APIs"]
    assert parsed["parsed_data"]["experience"][0]["title"] == "Senior Engineer"
    assert parsed["parsed_data"]["projects"][0]["name"] == "Candidate Platform"
    assert parsed["parsed_data"]["education"][0]["institution"] == "University of Testing"
    assert parsed["parsed_data"]["certifications"][0]["name"] == "Cloud Basics"
    assert parsed["parsed_data"]["skill_categories"][0]["category"] == "Backend"


def test_extract_text_routes_pdf_through_pdf_extractor(monkeypatch):
    """PDF file types should dispatch to the PDF extractor."""
    agent = ResumeParserAgent(enable_llm=False)

    monkeypatch.setattr(
        agent, "extract_text_from_pdf", lambda path: f"pdf text from {path}"
    )

    assert agent.extract_text("/tmp/resume.pdf", "pdf") == "pdf text from /tmp/resume.pdf"


def test_fallback_parse_produces_structured_sections():
    """Fallback parsing should preserve section structure with stable keys."""
    raw_text = """
    John Doe
    john.doe@example.com
    +1 (555) 123-4567
    Summary
    Backend engineer building APIs and AI tooling.
    Experience
    Senior Engineer Aug 2022 – Jan 2024
    Acme Corp New Delhi, India
    • Built backend APIs
    • Led parser design
    Projects
    Candidate Platform
    Python • FastAPI • PostgreSQL
    • Built resume parsing workflows
    Technical Skills
    Languages: Python, SQL
    Backend: FastAPI, PostgreSQL
    Certifications
    Cloud Basics – Example Academy (2024)
    """

    agent = ResumeParserAgent(enable_llm=False)
    structured = agent.fallback_structured_parse(raw_text)

    assert structured.summary == "Backend engineer building APIs and AI tooling."
    assert structured.schema_version == "resume.v2"
    assert structured.skills == ["Python", "SQL", "FastAPI", "PostgreSQL"]
    assert structured.skill_categories[0].category == "Languages"
    assert structured.experience[0].role_title == "Senior Engineer"
    assert structured.experience[0].company_name == "Acme Corp New Delhi, India"
    assert structured.experience[0].tasks_performed == ["Built backend APIs", "Led parser design"]
    assert structured.experience[0].title == "Senior Engineer"
    assert structured.projects[0].name == "Candidate Platform"
    assert structured.projects[0].technologies == ["Python", "FastAPI", "PostgreSQL"]
    assert structured.certifications[0].name == "Cloud Basics"


def test_fallback_parse_joins_wrapped_project_bullets_and_splits_certifications():
    """Wrapped PDF lines should stay inside the same project/certification entries."""
    raw_text = """
    Jane Doe
    jane@example.com
    Projects
    AI-Powered Video Summarizer with Gemini
    Python • Gemini 2.0 Flash API • Phidata
    • Built a multimodal video analysis application using Gemini 2.0 Flash for contextual summarization and key insight extraction from
    long-form video content.
    • Integrated an agentic workflow using Phidata tools and DuckDuckGo search to enhance AI-generated summaries with real-time web
    context and relevant external information.
    Certifications
    Data Science Masters 2.0 – PW Skills (2024) — Decode Python with DSA – PW Skills (2023)
    """

    agent = ResumeParserAgent(enable_llm=False)
    structured = agent.fallback_structured_parse(raw_text)

    assert len(structured.projects) == 1
    assert structured.projects[0].name == "AI-Powered Video Summarizer with Gemini"
    assert structured.projects[0].bullets == [
        "Built a multimodal video analysis application using Gemini 2.0 Flash for contextual summarization and key insight extraction from long-form video content.",
        "Integrated an agentic workflow using Phidata tools and DuckDuckGo search to enhance AI-generated summaries with real-time web context and relevant external information.",
    ]
    assert structured.certifications == [
        CertificationEntry(
            name="Data Science Masters 2.0",
            issuer="PW Skills",
            year="2024",
        ),
        CertificationEntry(
            name="Decode Python with DSA",
            issuer="PW Skills",
            year="2023",
        ),
    ]


def test_fallback_parse_separates_education_score_and_location():
    """Education parsing should keep score and location in their own fields."""
    raw_text = """
    Jane Doe
    Education
    GGSIPU, Delhi Technical Campus Jul 2021 – May 2025
    B.Tech – Computer Science and Engineering — CGPA: 8.2/10 Greater Noida, Delhi NCR
    """

    agent = ResumeParserAgent(enable_llm=False)
    structured = agent.fallback_structured_parse(raw_text)

    assert structured.education[0] == EducationEntry(
        institution="GGSIPU, Delhi Technical Campus",
        degree="B.Tech – Computer Science and Engineering",
        field="Computer Science and Engineering",
        location="Greater Noida, Delhi NCR",
        start_date="Jul 2021",
        end_date="May 2025",
        score="CGPA: 8.2/10",
    )
