"""API tests for resume upload, listing, detail, and deletion."""

from __future__ import annotations

from io import BytesIO
import zipfile

import pytest
import pytest_asyncio
from docx import Document
from fastapi import UploadFile
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.agents.resume_parser_agent import get_resume_parser_agent
from app.main import app
from app.models.job import Job
from app.models.resume import Resume
from app.routers import resumes as resumes_router
from app.services.resume_session_manager import ResumeSessionManager


class FakeResumeParserAgent:
    """Deterministic parser used to keep API tests local and fast."""

    async def parse_resume(self, file_path: str, file_type: str) -> dict:
        return {
            "candidate_name": "Jane Candidate",
            "phone_number": "+1 222 333 4444",
            "email": "jane@example.com",
            "raw_text": "Jane Candidate\njane@example.com\n+1 222 333 4444",
            "parsed_data": {
                "schema_version": "resume.v2",
                "summary": "Strong backend candidate.",
                "skills": ["Python", "FastAPI"],
                "skill_categories": [
                    {"category": "Backend", "items": ["Python", "FastAPI"]}
                ],
                "experience": [
                    {
                        "role_title": "Backend Engineer",
                        "company_name": "Example Corp",
                        "title": "Backend Engineer",
                        "company": "Example Corp",
                        "employment_type": None,
                        "location": None,
                        "from_date": "Jan 2024",
                        "to_date": "Present",
                        "start_date": "Jan 2024",
                        "end_date": "Present",
                        "is_current": True,
                        "duration_text": "Jan 2024 – Present",
                        "tasks_performed": ["Built APIs"],
                        "bullets": ["Built APIs"],
                    }
                ],
                "projects": [
                    {
                        "name": "Resume Parser",
                        "subtitle": None,
                        "technologies": ["Python", "FastAPI"],
                        "bullets": ["Built parser flows"],
                        "links": [],
                    }
                ],
                "education": [
                    {
                        "institution": "Example University",
                        "degree": "B.Sc Computer Science",
                        "field": None,
                        "location": None,
                        "start_date": None,
                        "end_date": None,
                        "score": None,
                    }
                ],
                "certifications": [],
                "links": [],
            },
        }


class OneSuccessThenFailureParser:
    """Parser that persists one resume, then fails the next file."""

    def __init__(self):
        self.calls = 0

    async def parse_resume(self, file_path: str, file_type: str) -> dict:
        self.calls += 1
        if self.calls > 1:
            raise RuntimeError("parse failed")
        return {
            "candidate_name": "Saved Candidate",
            "phone_number": "+1 111 222 3333",
            "email": "saved@example.com",
            "raw_text": "Saved Candidate\nsaved@example.com",
            "parsed_data": {"schema_version": "resume.v2"},
        }


@pytest_asyncio.fixture
async def job_for_resume(db_session, test_user) -> Job:
    """Create a job record that can own uploaded resumes."""
    job = Job(
        user_id=test_user.id,
        title="Backend Engineer",
        description="Build APIs",
        requirements="FastAPI, PostgreSQL",
        status="active",
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)
    return job


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Generate an in-memory DOCX resume."""
    document = Document()
    document.add_paragraph("Jane Candidate")
    document.add_paragraph("jane@example.com")
    document.add_paragraph("+1 222 333 4444")
    document.add_paragraph("Skills")
    document.add_paragraph("Python")
    document.add_paragraph("FastAPI")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest_asyncio.fixture(autouse=True)
async def override_resume_parser():
    """Use a local fake parser for resume endpoint tests."""
    app.dependency_overrides[get_resume_parser_agent] = lambda: FakeResumeParserAgent()
    yield
    app.dependency_overrides.pop(get_resume_parser_agent, None)


@pytest.mark.asyncio
async def test_upload_resume(
    authenticated_client: AsyncClient, job_for_resume: Job, sample_docx_bytes: bytes
):
    """Uploading a DOCX resume should create a parsed resume record."""
    response = await authenticated_client.post(
        f"/api/jobs/{job_for_resume.id}/resumes",
        files=[("files", ("candidate.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )

    assert response.status_code == 201
    data = response.json()
    assert len(data) == 1
    assert data[0]["candidate_name"] == "Jane Candidate"
    assert data[0]["status"] == "parsed"
    assert data[0]["file_type"] == "docx"
    assert data[0]["parsed_data"]["schema_version"] == "resume.v2"
    assert data[0]["parsed_data"]["skills"] == ["Python", "FastAPI"]
    assert data[0]["parsed_data"]["experience"][0]["role_title"] == "Backend Engineer"
    assert data[0]["parsed_data"]["projects"][0]["name"] == "Resume Parser"


@pytest.mark.asyncio
async def test_list_resumes_for_job(
    authenticated_client: AsyncClient, job_for_resume: Job, sample_docx_bytes: bytes
):
    """Listing resumes should return the uploaded records for the job."""
    await authenticated_client.post(
        f"/api/jobs/{job_for_resume.id}/resumes",
        files=[("files", ("candidate.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )

    response = await authenticated_client.get(f"/api/jobs/{job_for_resume.id}/resumes")
    assert response.status_code == 200
    data = response.json()
    assert len(data) == 1
    assert data[0]["email"] == "jane@example.com"


@pytest.mark.asyncio
async def test_get_resume_detail(
    authenticated_client: AsyncClient, job_for_resume: Job, sample_docx_bytes: bytes
):
    """Resume detail endpoint should include raw text."""
    upload_response = await authenticated_client.post(
        f"/api/jobs/{job_for_resume.id}/resumes",
        files=[("files", ("candidate.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    resume_id = upload_response.json()[0]["id"]

    response = await authenticated_client.get(f"/api/resumes/{resume_id}")
    assert response.status_code == 200
    data = response.json()
    assert data["candidate_name"] == "Jane Candidate"
    assert "Jane Candidate" in data["raw_text"]


@pytest.mark.asyncio
async def test_delete_resume(
    authenticated_client: AsyncClient, job_for_resume: Job, sample_docx_bytes: bytes
):
    """Deleting a resume should remove it from the list."""
    upload_response = await authenticated_client.post(
        f"/api/jobs/{job_for_resume.id}/resumes",
        files=[("files", ("candidate.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    resume_id = upload_response.json()[0]["id"]

    delete_response = await authenticated_client.delete(f"/api/resumes/{resume_id}")
    assert delete_response.status_code == 204

    list_response = await authenticated_client.get(f"/api/jobs/{job_for_resume.id}/resumes")
    assert list_response.status_code == 200
    assert list_response.json() == []


@pytest.mark.asyncio
async def test_upload_resume_rejects_invalid_file_type(
    authenticated_client: AsyncClient, job_for_resume: Job
):
    """Only PDF and DOCX uploads should be accepted."""
    response = await authenticated_client.post(
        f"/api/jobs/{job_for_resume.id}/resumes",
        files=[("files", ("candidate.txt", b"plain text", "text/plain"))],
    )

    assert response.status_code == 422
    assert "Only PDF and DOCX" in response.json()["detail"]


@pytest.mark.asyncio
async def test_cannot_access_resume_for_other_users_job(
    client: AsyncClient,
    authenticated_client: AsyncClient,
    db_session,
    test_user,
    sample_docx_bytes: bytes,
):
    """Resume endpoints should be scoped to the owner of the parent job."""
    from app.core.security import create_access_token, hash_password
    from app.models.user import User

    job = Job(
        user_id=test_user.id,
        title="Private Role",
        description="Hidden",
        requirements=None,
        status="active",
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    upload_response = await authenticated_client.post(
        f"/api/jobs/{job.id}/resumes",
        files=[("files", ("candidate.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    resume_id = upload_response.json()[0]["id"]

    user2 = User(
        email="resume-other@example.com",
        hashed_password=hash_password("Password123"),
        full_name="Other Resume User",
        is_active=True,
    )
    db_session.add(user2)
    await db_session.commit()
    await db_session.refresh(user2)

    token2 = create_access_token({"sub": str(user2.id)})
    client.headers.update({"Authorization": f"Bearer {token2}"})

    response = await client.get(f"/api/resumes/{resume_id}")
    assert response.status_code == 404


@pytest.mark.asyncio
async def test_background_processing_commits_each_successful_resume(
    monkeypatch, db_session: AsyncSession, job_for_resume: Job
):
    """A later file failure must not hide already processed candidates from refreshes."""
    session_factory = async_sessionmaker(
        db_session.bind, class_=AsyncSession, expire_on_commit=False
    )
    monkeypatch.setattr("app.database.async_session_factory", session_factory)

    session_manager = ResumeSessionManager(session_ttl_seconds=3600)
    session_id = session_manager.create_session(
        job_id=str(job_for_resume.id),
        user_id=str(job_for_resume.user_id),
        total_resumes=2,
    )

    await resumes_router._process_resumes_background(
        session_id=session_id,
        job_id=job_for_resume.id,
        file_paths=[
            ("uploads/one.pdf", "one.pdf", "pdf"),
            ("uploads/two.pdf", "two.pdf", "pdf"),
        ],
        job_description=None,
        session_manager=session_manager,
        storage=object(),
        parser=OneSuccessThenFailureParser(),
    )

    result = await db_session.execute(select(Resume).where(Resume.job_id == job_for_resume.id))
    saved_resumes = result.scalars().all()
    assert len(saved_resumes) == 1
    assert saved_resumes[0].candidate_name == "Saved Candidate"

    events = []
    queue = session_manager.get_session(session_id).event_queue
    while not queue.empty():
        events.append(queue.get_nowait())

    completed_events = [event for event in events if event.event_type == "completed"]
    assert completed_events
    assert completed_events[0].resume_id == str(saved_resumes[0].id)
    assert any(event.event_type == "error" for event in events)


@pytest.mark.asyncio
async def test_export_resumes_xlsx_skips_candidates_without_contact_info(
    authenticated_client: AsyncClient, db_session: AsyncSession, job_for_resume: Job
):
    """Candidate export should include all contactable candidates and AI fields."""
    db_session.add_all(
        [
            Resume(
                job_id=job_for_resume.id,
                candidate_name="Contactable Candidate",
                phone_number="+1 222 333 4444",
                email="contact@example.com",
                file_path=f"uploads/resumes/{job_for_resume.id}/67dda1c5-1ecb-4e6c-b4eb-5d063f5436e5-candidate.pdf",
                file_type="pdf",
                matching_score=91.2,
                match_explanation="Strong FastAPI and PostgreSQL background.",
                status="parsed",
            ),
            Resume(
                job_id=job_for_resume.id,
                candidate_name="No Contact Candidate",
                phone_number=None,
                email=None,
                file_path=f"uploads/resumes/{job_for_resume.id}/no-contact.pdf",
                file_type="pdf",
                matching_score=80,
                match_explanation="Should not export.",
                status="parsed",
            ),
        ]
    )
    await db_session.commit()

    response = await authenticated_client.get(f"/api/jobs/{job_for_resume.id}/resumes/export")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "candidates" in response.headers["content-disposition"]

    with zipfile.ZipFile(BytesIO(response.content)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode()

    assert "Candidate Name" in sheet_xml
    assert "Contactable Candidate" in sheet_xml
    assert "contact@example.com" in sheet_xml
    assert "candidate.pdf" in sheet_xml
    assert "91%" in sheet_xml
    assert "Strong FastAPI and PostgreSQL background." in sheet_xml
    assert "No Contact Candidate" not in sheet_xml
